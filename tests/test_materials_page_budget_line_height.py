"""The pre-render page budget must measure lines the way LibreOffice draws them.

The gate exists to catch one-page overflow *before* LibreOffice runs.  A model
that under-measures still passes every document that actually overflows, which
is the exact failure it was added to prevent: the character-count estimator
``estimate_canonical_capacity`` reports ratio 1.0 on a document that renders as
two pages, and the overflow is only discovered after the PDF exists.

Two reference facts, measured rather than fitted, anchor these tests.  On a
python-docx default template (docDefaults ``w:sz="22"``, ``w:after="200"``,
``w:line="276"``) with a 648pt body, one paragraph of 11pt Caladea at 1.15x
occupies 24.8pt of pitch -- a 14.8pt line plus the 10pt space-after -- and 26
identical paragraphs fit one page while 27 spill onto a second.

An earlier model computed 21.5pt for that same paragraph: it resolved the font
at a hard-coded 10pt because nothing read the docDefaults ``rPrDefault``, and
it took ``max(multiple, leading)`` where LibreOffice multiplies the font's
natural line height by the multiple.  It therefore flipped at 31 paragraphs and
let four paragraphs of real overflow reach the PDF stage.
"""

from __future__ import annotations

import pytest

from tools.workflow.materials_renderer import (
    _page_budget_report,
    _paragraph_points,
    _resolve_font_pt,
    estimate_canonical_capacity,
)

_TEXT = "A tailored evidence line for the current job requirement."
_BODY_PT = 648.0
_FLIP_COUNT = 26
# 26 * 24.835 = 645.7 (one page); 27 * 24.835 = 670.5 (two pages).
_FLIP_POINTS = {26: 645.7, 27: 670.5}


def _doc(count: int):
    from docx import Document
    from docx.shared import Pt

    document = Document()
    section = document.sections[0]
    section.page_height = Pt(792)
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    for _ in range(count):
        document.add_paragraph(_TEXT)
    return document


def test_doc_defaults_supply_the_font_size_the_renderer_actually_uses():
    document = _doc(1)
    # The default template declares 11pt only in docDefaults rPrDefault; the
    # style chain carries no size, so a resolver that stops at the chain falls
    # back to a hard-coded 10pt and under-measures every line by 9%.
    assert _resolve_font_pt(document.paragraphs[0], document) == pytest.approx(11.0, abs=0.01)


def test_paragraph_points_match_the_measured_libreoffice_pitch():
    document = _doc(1)
    points = _paragraph_points(document.paragraphs[0], material="cv", document=document)
    assert points == pytest.approx(24.8, abs=0.5)


def test_gate_measures_a_full_page_the_way_libreoffice_renders_it():
    for count, expected in _FLIP_POINTS.items():
        report = _page_budget_report(_doc(count), material="cv")
        assert report["budget_points"] == pytest.approx(_BODY_PT, abs=0.1)
        assert report["points"] == pytest.approx(expected, abs=1.5)


def test_gate_flips_where_libreoffice_flips():
    assert _page_budget_report(_doc(_FLIP_COUNT), material="cv")["over_by_points"] <= 0
    assert _page_budget_report(_doc(_FLIP_COUNT + 1), material="cv")["over_by_points"] > 0


def test_gate_catches_overflow_the_character_estimator_misses():
    """The task's completion criterion: capacity says fine, geometry says stop.

    A canonical that is exactly as long as its master reports ratio 1.0 to the
    cheap character estimator, yet occupies more vertical space than the page
    because every paragraph carries the template's default spacing.
    """

    blocks = [{"source_style": "Normal", "text": _TEXT} for _ in range(_FLIP_COUNT + 1)]
    estimate = estimate_canonical_capacity(
        {"cv": {"blocks": blocks}, "cover_letter": {"blocks": []}},
        {"cv": {"blocks": blocks}, "cover_letter": {"blocks": []}},
    )["cv"]
    assert estimate["ratio"] == 1.0
    assert estimate["over_master_lines"] == 0
    report = _page_budget_report(_doc(_FLIP_COUNT + 1), material="cv")
    assert report["over_by_points"] > 0
    assert report["material"] == "cv"
