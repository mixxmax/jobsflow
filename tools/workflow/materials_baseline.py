"""Lane-master content baselines for bounded CV/CL tailoring.

The selected lane masters are both the visual templates and the semantic
starting point.  A drafting model receives their ordered text blocks and may
submit only bounded rewrites, reordering, merges, or additions.  Blocks that
the model does not mention are retained by the host.
"""

from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from datetime import datetime, timezone
from math import ceil
from pathlib import Path
from typing import Any

from tools.io_utils import atomic_write_json
from tools.job_materials.paths import find_latest_cl_master_docx, find_latest_master_docx
from tools.workflow.materials_hashes import container_hash

BASELINE_NAME = "materials_baseline.json"
BASELINE_SCHEMA_VERSION = 1
# Bump when block extraction, section classification or host-managed marking
# changes.  The cache predicate requires this exact version, so a parser
# change rebuilds every package baseline instead of silently reusing blocks
# parsed by the previous extractor.
BASELINE_EXTRACTOR_VERSION = 2
_PLACEHOLDER_RE = re.compile(r"\[[^\]]+\]|\{[^}]+\}|\b(?:TBD|TODO)\b", re.I)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _digest(value: Any) -> str:
    payload = dict(value or {}) if isinstance(value, dict) else value
    if isinstance(payload, dict):
        payload = {key: item for key, item in payload.items() if key not in {"created_at", "baseline_sha256"}}
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def baseline_path(package: Path) -> Path:
    return Path(package) / BASELINE_NAME


def load_content_baseline(package: Path) -> dict[str, Any]:
    package = Path(package)
    candidates = [
        baseline_path(package),
        # vNext owns the frozen baseline for the current job.  Treat it as
        # the compatibility projection instead of recompiling a second,
        # potentially divergent baseline for the legacy renderer/helpers.
        package / "materials_vnext" / "baseline_snapshot.json",
    ]
    try:
        for candidate in candidates:
            if not candidate.is_file():
                continue
            value = json.loads(candidate.read_text(encoding="utf-8"))
            if isinstance(value, dict):
                return value
    except (OSError, json.JSONDecodeError):
        pass
    return {}


def content_baseline_digest(baseline_or_package: dict[str, Any] | Path) -> str:
    baseline = load_content_baseline(baseline_or_package) if isinstance(baseline_or_package, Path) else baseline_or_package
    return _digest(baseline)


def _section_name(text: str, current: str) -> str:
    folded = " ".join(str(text or "").casefold().split())
    if "summary" in folded or "profile" in folded:
        return "summary"
    if "core" in folded or "expertise" in folded or "competenc" in folded:
        return "core"
    if "experience" in folded or "employment" in folded:
        return "experience"
    if "education" in folded or "academic" in folded or "degree" in folded or "教育" in folded or "学位" in folded:
        return "education"
    if "qualification" in folded or "language" in folded or "admission" in folded or "certification" in folded or "licen" in folded or "资格" in folded or "语言" in folded or "执业" in folded:
        return "qualifications"
    return current


def _master_blocks(path: Path, *, material: str) -> list[dict[str, Any]]:
    from docx import Document

    document = Document(str(path))
    output: list[dict[str, Any]] = []
    section = "header" if material == "cv" else "letter"
    experience_id = ""
    normal_seen = 0
    for paragraph in document.paragraphs:
        text = " ".join(str(paragraph.text or "").split()).strip()
        if not text:
            continue
        style = str(paragraph.style.name or "Normal")
        position = len(output) + 1
        if material == "cv":
            if style == "Resume Section":
                section = _section_name(text, section)
                block_type = "heading"
                experience_id = ""
            elif style == "Job Heading":
                section = "experience"
                experience_id = f"base-experience-{position:03d}"
                block_type = "heading"
            elif style == "Resume Bullet":
                block_type = "bullet"
            elif style == "Compact Line":
                block_type = "bullet" if section in {"core", "experience"} else "paragraph"
            elif normal_seen < 2:
                block_type = "contact"
                section = "contact"
                normal_seen += 1
            else:
                block_type = "paragraph"
        else:
            if normal_seen < 2 and style == "Normal":
                block_type = "contact"
                section = "contact"
                normal_seen += 1
            elif re.fullmatch(r"(?:Hiring Manager|\[\s*Company(?:\s+address)?\s*\])", text, re.I):
                # Recipient identity is host-managed.  It is material-facing
                # output, but it is not a drafting decision: a disclosed
                # employer is substituted by the host, while an undisclosed
                # recruiter client receives a neutral placeholder.  Marking
                # these blocks prevents a model from deleting or rewriting
                # the company line after looking at another package.
                block_type = "paragraph"
                section = "recipient"
            elif text.casefold().startswith("re:"):
                block_type = "heading"
                section = "subject"
            elif text.casefold().startswith(("yours sincerely", "kind regards", "best regards")):
                block_type = "signoff"
                section = "closing"
            elif style == "Letter Bullet":
                block_type = "bullet"
                section = "evidence"
            else:
                block_type = "paragraph"
                section = "body"
        block_id = f"base-{'cv' if material == 'cv' else 'cl'}-{position:03d}"
        optional_slot = bool(
            material == "cover_letter"
            and re.fullmatch(r"\[\s*company\s+address\s*\]", text, re.I)
        )
        output.append(
            {
                "id": block_id,
                "type": block_type,
                "text": text,
                "section": section,
                "experience_id": experience_id,
                "priority": position,
                "source_style": style,
                "requires_rewrite": bool(_PLACEHOLDER_RE.search(text)),
                "content_floor": not optional_slot,
                "host_managed_optional": optional_slot,
                "host_managed": bool(
                    material == "cover_letter"
                    and section in {"contact", "date", "recipient", "subject"}
                ),
            }
        )
    if not output:
        raise ValueError(f"base_content_missing:{material}")
    return output


def build_content_baseline(*, workspace: Path, package: Path, lane: str) -> dict[str, Any]:
    """Extract and freeze the selected lane masters as a package baseline."""

    workspace = Path(workspace)
    package = Path(package)
    cv_path = find_latest_master_docx(str(lane or "").upper(), workspace)
    cl_path = find_latest_cl_master_docx(str(lane or "").upper(), workspace)
    if cv_path is None:
        raise ValueError("base_template_missing:cv")
    if cl_path is None:
        raise ValueError("base_template_missing:cover_letter")
    source_hashes = {"cv": container_hash(cv_path), "cover_letter": container_hash(cl_path)}
    current = load_content_baseline(package)
    if (
        current
        and current.get("source_hashes") == source_hashes
        and current.get("baseline_sha256") == _digest(current)
        and int(current.get("extractor_version") or 0) == BASELINE_EXTRACTOR_VERSION
    ):
        return current
    baseline: dict[str, Any] = {
        "schema_version": BASELINE_SCHEMA_VERSION,
        "extractor_version": BASELINE_EXTRACTOR_VERSION,
        "artifact_type": "jobsflow_lane_content_baseline",
        "lane": str(lane or "").upper(),
        "source_paths": {"cv": str(cv_path.resolve()), "cover_letter": str(cl_path.resolve())},
        "source_hashes": source_hashes,
        "contract": {
            "mode": "bounded_incremental_transform",
            "unmentioned_blocks": "retain",
            "deletion_allowed": False,
            "host_managed_identity": "CV/CL contact, recipient, role and company identity lines are generated from the current job contract; drafting models may not rewrite or delete them.",
            "allowed_actions": ["retain", "rewrite", "reorder", "merge", "add"],
            "content_floor": "Every evidence-bearing baseline block must remain represented; host-managed optional address slots are excluded, and tailoring may rewrite, reorder or merge but may not silently delete evidence.",
        },
        "cv": {"blocks": _master_blocks(cv_path, material="cv")},
        "cover_letter": {"blocks": _master_blocks(cl_path, material="cover_letter")},
        "created_at": _now(),
    }
    baseline["baseline_sha256"] = _digest(baseline)
    atomic_write_json(baseline_path(package), baseline)
    return baseline


def baseline_task_view(baseline: dict[str, Any]) -> dict[str, Any]:
    """Return model-visible content without leaking local master paths."""

    return {
        "schema_version": baseline.get("schema_version"),
        "artifact_type": baseline.get("artifact_type"),
        "lane": baseline.get("lane"),
        "baseline_sha256": baseline.get("baseline_sha256"),
        "contract": dict(baseline.get("contract") or {}),
        "cv": {"blocks": list((baseline.get("cv") or {}).get("blocks") or [])},
        "cover_letter": {"blocks": list((baseline.get("cover_letter") or {}).get("blocks") or [])},
    }


def plan_jd_anchor_catalog(plan: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Normalize plan duties/requirements into stable transform anchor IDs."""

    plan = dict(plan or {})
    output: list[dict[str, Any]] = []
    raw_anchors = plan.get("jd_anchors") or []
    for raw in raw_anchors if isinstance(raw_anchors, list) else []:
        if isinstance(raw, dict):
            text = " ".join(str(raw.get("text") or raw.get("label") or "").split()).strip()
            anchor_id = str(raw.get("id") or "").strip()
            priority = raw.get("priority")
        else:
            text = " ".join(str(raw or "").split()).strip()
            anchor_id = ""
            priority = None
        if text:
            output.append(
                {
                    "id": anchor_id or f"JD-{len(output) + 1:03d}",
                    "text": text,
                    "priority": priority if priority is not None else len(output) + 1,
                }
            )
    if not output:
        seen: set[str] = set()
        for field in ("duties", "requirements", "themes"):
            values = plan.get(field) or []
            if not isinstance(values, list):
                continue
            for raw in values:
                text = " ".join(str(raw or "").split()).strip()
                key = text.casefold()
                if not text or key in seen:
                    continue
                seen.add(key)
                output.append(
                    {
                        "id": f"JD-{len(output) + 1:03d}",
                        "text": text,
                        "priority": len(output) + 1,
                    }
                )
    return output or [
        {
            "id": "JD-001",
            "text": "the selected JD duties and requirements",
            "priority": 1,
        }
    ]


def baseline_transform_task_schema(
    baseline: dict[str, Any],
    *,
    job_id: str,
    jd_anchors: list[dict[str, Any]] | None = None,
    contract: str = "legacy",
) -> dict[str, Any]:
    """Return the single model-facing drafting contract.

    The canonical document is an internal host artifact. Drafting models only
    submit this bounded delta, so a weaker model cannot accidentally select a
    full-document replacement path advertised by an older response schema.
    """

    if str(contract or "").casefold() == "vnext":
        # The public gateway has one model-facing vocabulary.  Keep this
        # explicit instead of letting the compatibility module advertise its
        # older ``baseline_id/merge/add`` shape to a vNext caller.
        return {
            "schema_version": 1,
            "artifact_type": "jobsflow_baseline_transform",
            "job_id": str(job_id),
            "baseline_sha256": str(baseline.get("baseline_sha256") or ""),
            "jd_anchor_catalog": list(jd_anchors or []),
            "unmentioned_blocks": "retain",
            "deletion_allowed": False,
            "allowed_actions": ["replace", "append_after", "reorder"],
            "scope_policy": {
                "normal": "prefer a focused set of high-value JD changes",
                "broad_review": "more than roughly 35% of a document's baseline blocks routes to a stronger auditor",
                "full_replacement_equivalent": "more than 60% of baseline blocks is rejected",
            },
            "host_managed_blocks": [
                {
                    "material": material,
                    "baseline_id": str(block.get("id") or ""),
                    "section": str(block.get("section") or ""),
                    "text": str(block.get("text") or ""),
                }
                for material in ("cv", "cover_letter")
                for block in (baseline.get(material) or {}).get("blocks") or []
                if isinstance(block, dict) and block.get("host_managed")
            ],
            "operations": (
                "array of {material, action: replace|append_after|reorder, "
                "target_id, before_text/after_text or block, jd_anchor_ids}"
            ),
            "instruction": (
                "Use the lane baseline as the content master. Return only the bounded "
                "JD-specific delta in operations; do not rebuild or silently shorten "
                "the CV or Cover Letter."
            ),
        }

    # Compatibility contract for the retired adapter.  It remains available
    # for migration/old fixture callers, but the public workflow never emits
    # it as the vNext model task packet.
    return {
        "schema_version": 1,
        "artifact_type": "jobsflow_baseline_transform",
        "job_id": str(job_id),
        "baseline_sha256": str(baseline.get("baseline_sha256") or ""),
        "jd_anchor_catalog": list(jd_anchors or []),
        "unmentioned_blocks": "retain",
        "deletion_allowed": False,
        "allowed_actions": ["rewrite", "reorder", "merge", "add"],
        "scope_policy": {
            "normal": "prefer a focused set of high-value JD changes",
            "broad_review": "more than roughly 35% of a document's baseline blocks receives extra review",
            "full_replacement_equivalent": "more than 60% of baseline blocks is rejected",
        },
        "host_managed_blocks": [
            {
                "material": material,
                "baseline_id": str(block.get("id") or ""),
                "section": str(block.get("section") or ""),
                "text": str(block.get("text") or ""),
            }
            for material in ("cv", "cover_letter")
            for block in (baseline.get(material) or {}).get("blocks") or []
            if isinstance(block, dict) and block.get("host_managed")
        ],
        "changes": "array of JD-anchored rewrite/reorder/merge operations",
        "additions": "array of truthful JD-anchored blocks",
        "instruction": (
            "Use the lane baseline as the content master. Return only the bounded "
            "JD-specific delta; wording may change materially when useful, but do "
            "not rebuild or silently shorten the CV or Cover Letter."
        ),
    }


def _host_substitutions(
    text: str,
    context: dict[str, Any],
    plan: dict[str, Any] | None = None,
) -> str:
    """Fill deterministic identity/date tokens before model tailoring."""

    role = str(context.get("role_primary") or context.get("role") or "").strip()
    employer = str(context.get("employer_name") or context.get("company_out") or "").strip()
    candidate = str(context.get("candidate_name") or "").strip()
    application_date = str(context.get("application_date") or datetime.now(timezone.utc).strftime("%d %B %Y"))
    plan = dict(plan or {})
    primary_anchor = next(
        (
            str(item).strip()
            for field in ("duties", "requirements", "themes")
            for item in (plan.get(field) or [])
            if str(item).strip()
        ),
        role or "the role's principal responsibilities",
    )
    replacements = {
        "date": application_date,
        "role": role,
        "position": role,
        "company": employer or "the hiring organisation",
        "employer": employer or "the hiring organisation",
        "candidate": candidate,
        "candidate name": candidate,
        "your name": candidate,
        "name": candidate,
    }

    def replace(match: re.Match[str]) -> str:
        key = " ".join(re.sub(r"[_-]+", " ", match.group(1).casefold()).split())
        if key in replacements:
            return replacements[key] or match.group(0)
        # Candidate placeholders are setup blockers, not JD-tailoring slots.
        # Leaving them intact makes the canonical gate fail visibly instead of
        # replacing a missing personal fact with unrelated job wording.
        if key.startswith("your ") or "contact" in key:
            return match.group(0)
        return primary_anchor

    return re.sub(r"\[([^\]]+)\]", replace, str(text or ""))


def canonical_from_baseline(
    baseline: dict[str, Any],
    *,
    job_id: str,
    context: dict[str, Any],
    plan: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Create the safe seed; every evidence-bearing master block is retained."""

    if not baseline or baseline.get("artifact_type") != "jobsflow_lane_content_baseline":
        raise ValueError("content_baseline_missing")
    dispositions: dict[str, dict[str, Any]] = {}
    documents: dict[str, Any] = {}
    for material in ("cv", "cover_letter"):
        blocks: list[dict[str, Any]] = []
        for raw in (baseline.get(material) or {}).get("blocks") or []:
            if not isinstance(raw, dict) or not raw.get("id"):
                continue
            if raw.get("host_managed_optional") and not raw.get("content_floor", True):
                continue
            block = {
                key: deepcopy(value)
                for key, value in raw.items()
                if key not in {"source_style", "requires_rewrite", "content_floor"}
            }
            block["text"] = _host_substitutions(
                str(block.get("text") or ""), context, plan
            )
            block["jd_anchor_ids"] = []
            block["baseline_refs"] = [str(raw["id"])]
            block["change_action"] = "retain"
            blocks.append(block)
            dispositions[str(raw["id"])] = {
                "material": material,
                "action": "retain",
                "target_id": str(block["id"]),
            }
        documents[material] = {"blocks": blocks}
    plan = dict(plan or {})
    return {
        "schema_version": 1,
        "artifact_type": "jobsflow_canonical_cv_cl",
        "job_id": str(job_id),
        "compiled_from": "lane_content_baseline",
        "baseline_sha256": str(baseline.get("baseline_sha256") or content_baseline_digest(baseline)),
        "baseline_dispositions": dispositions,
        "coverage_dispositions": dict(plan.get("coverage_dispositions") or {}),
        "jd_anchors": plan_jd_anchor_catalog(plan),
        **documents,
    }


def apply_baseline_transform(
    baseline: dict[str, Any],
    transform: dict[str, Any],
    *,
    job_id: str,
    context: dict[str, Any],
    plan: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Apply a small model delta; omitted baseline blocks stay untouched."""

    if not isinstance(transform, dict) or transform.get("artifact_type") != "jobsflow_baseline_transform":
        raise ValueError("baseline_transform_required")
    if str(transform.get("job_id") or "") != str(job_id):
        raise ValueError("baseline_transform_job_id_mismatch")
    expected_digest = str(baseline.get("baseline_sha256") or content_baseline_digest(baseline))
    if str(transform.get("baseline_sha256") or "") != expected_digest:
        raise ValueError("baseline_transform_stale")
    draft = canonical_from_baseline(baseline, job_id=job_id, context=context, plan=plan)
    allowed_anchor_ids = {
        str(item.get("id") or "")
        for item in plan_jd_anchor_catalog(plan)
        if isinstance(item, dict) and str(item.get("id") or "")
    }
    index: dict[str, tuple[str, dict[str, Any]]] = {}
    for material in ("cv", "cover_letter"):
        for block in draft[material]["blocks"]:
            for baseline_id in block.get("baseline_refs") or []:
                index[str(baseline_id)] = (material, block)

    changes = transform.get("changes") or []
    if not isinstance(changes, list):
        raise ValueError("baseline_transform_changes_not_list")
    changed_ids: set[str] = set()
    tailored_materials: set[str] = set()
    for raw in changes:
        if not isinstance(raw, dict):
            raise ValueError("baseline_transform_change_not_object")
        action = str(raw.get("action") or "").casefold()
        anchors = raw.get("jd_anchor_ids") or []
        if not isinstance(anchors, list) or not any(str(item).strip() for item in anchors):
            raise ValueError("baseline_transform_jd_anchor_missing")
        if any(str(item).strip() not in allowed_anchor_ids for item in anchors):
            raise ValueError("baseline_transform_jd_anchor_unknown")
        if action in {"delete", "omit", "remove"}:
            raise ValueError("baseline_deletion_forbidden")
        if action == "merge":
            material = str(raw.get("material") or "").casefold()
            baseline_ids = raw.get("baseline_ids") or []
            if not isinstance(baseline_ids, list):
                raise ValueError("baseline_transform_merge_ids_not_list")
            baseline_ids = [str(item) for item in baseline_ids if str(item).strip()]
            if len(baseline_ids) < 2 or len(set(baseline_ids)) != len(baseline_ids):
                raise ValueError("baseline_transform_merge_ids_invalid")
            targets = [index.get(item) for item in baseline_ids]
            if any(target is None or target[0] != material for target in targets):
                raise ValueError("baseline_transform_merge_target_invalid")
            if any(bool(target[1].get("host_managed")) for target in targets if target is not None):
                raise ValueError("baseline_transform_host_managed_block")
            if any(item in changed_ids for item in baseline_ids):
                raise ValueError("baseline_transform_merge_target_duplicate")
            text = " ".join(str(raw.get("text") or "").split()).strip()
            if not text:
                raise ValueError("baseline_transform_merge_text_missing")
            blocks = draft[material]["blocks"]
            source_blocks = [target[1] for target in targets if target is not None]
            position = min(blocks.index(block) for block in source_blocks)
            first = source_blocks[0]
            for block in source_blocks:
                blocks.remove(block)
            merged_id = f"merge-{baseline_ids[0]}"
            merged = {
                "id": merged_id,
                "type": str(raw.get("type") or first.get("type") or "paragraph"),
                "text": text,
                "section": str(raw.get("section") or first.get("section") or ""),
                "experience_id": str(raw.get("experience_id") or first.get("experience_id") or ""),
                "priority": raw.get("priority") if raw.get("priority") is not None else first.get("priority", position),
                "jd_anchor_ids": [str(item) for item in (raw.get("jd_anchor_ids") or []) if str(item).strip()],
                "baseline_refs": baseline_ids,
                "change_action": "merge",
            }
            blocks.insert(position, merged)
            for baseline_id in baseline_ids:
                draft["baseline_dispositions"][baseline_id] = {
                    "material": material,
                    "action": "merge",
                    "target_id": merged_id,
                }
            changed_ids.update(baseline_ids)
            tailored_materials.add(material)
            continue
        if action not in {"rewrite", "reorder"}:
            raise ValueError(f"baseline_transform_action_invalid:{action or 'missing'}")
        baseline_id = str(raw.get("baseline_id") or "")
        material = str(raw.get("material") or "").casefold()
        target = index.get(baseline_id)
        if not target or target[0] != material:
            raise ValueError(f"baseline_transform_target_invalid:{baseline_id}")
        if bool(target[1].get("host_managed")):
            raise ValueError(f"baseline_transform_host_managed_block:{baseline_id}")
        if baseline_id in changed_ids:
            raise ValueError(f"baseline_transform_target_duplicate:{baseline_id}")
        block = target[1]
        if action == "rewrite":
            text = " ".join(str(raw.get("text") or "").split()).strip()
            if not text:
                raise ValueError(f"baseline_transform_text_missing:{baseline_id}")
            block["text"] = text
        else:
            try:
                position = int(raw.get("position"))
            except (TypeError, ValueError):
                raise ValueError(f"baseline_transform_position_missing:{baseline_id}") from None
            blocks = draft[material]["blocks"]
            if position < 0 or position >= len(blocks):
                raise ValueError(f"baseline_transform_position_invalid:{baseline_id}")
            blocks.remove(block)
            blocks.insert(position, block)
        block["change_action"] = action
        block["jd_anchor_ids"] = [str(item) for item in anchors if str(item).strip()] if isinstance(anchors, list) else []
        if raw.get("priority") is not None:
            block["priority"] = raw.get("priority")
        draft["baseline_dispositions"][baseline_id] = {
            "material": material,
            "action": action,
            "target_id": str(block["id"]),
        }
        changed_ids.add(baseline_id)
        tailored_materials.add(material)

    additions = transform.get("additions") or []
    if not isinstance(additions, list):
        raise ValueError("baseline_transform_additions_not_list")
    if not changes and not additions:
        raise ValueError("baseline_transform_empty")
    for number, raw in enumerate(additions, start=1):
        if not isinstance(raw, dict):
            raise ValueError("baseline_transform_addition_not_object")
        material = str(raw.get("material") or "").casefold()
        if material not in {"cv", "cover_letter"}:
            raise ValueError("baseline_transform_addition_material_invalid")
        text = " ".join(str(raw.get("text") or "").split()).strip()
        if not text:
            raise ValueError("baseline_transform_addition_text_missing")
        anchors = raw.get("jd_anchor_ids") or []
        if not isinstance(anchors, list) or not any(str(item).strip() for item in anchors):
            raise ValueError("baseline_transform_jd_anchor_missing")
        if any(str(item).strip() not in allowed_anchor_ids for item in anchors):
            raise ValueError("baseline_transform_jd_anchor_unknown")
        block_type = str(raw.get("type") or ("bullet" if material == "cv" else "paragraph"))
        if block_type not in {"heading", "contact", "paragraph", "bullet", "signoff"}:
            raise ValueError("baseline_transform_addition_type_invalid")
        block = {
            "id": f"add-{'cv' if material == 'cv' else 'cl'}-{number:03d}",
            "type": block_type,
            "text": text,
            "section": str(raw.get("section") or ("experience" if material == "cv" else "body")),
            "experience_id": str(raw.get("experience_id") or ""),
            "priority": raw.get("priority") if raw.get("priority") is not None else 999,
            "jd_anchor_ids": [str(item) for item in anchors if str(item).strip()],
            "baseline_refs": [],
            "change_action": "add",
        }
        target_blocks = draft[material]["blocks"]
        insert_after = str(raw.get("insert_after") or "")
        position = next((i + 1 for i, item in enumerate(target_blocks) if str(item.get("id")) == insert_after), len(target_blocks))
        target_blocks.insert(position, block)
        tailored_materials.add(material)
    for missing_material in sorted({"cv", "cover_letter"} - tailored_materials):
        raise ValueError(f"baseline_transform_material_missing:{missing_material}")
    scope: dict[str, Any] = {}
    for material in ("cv", "cover_letter"):
        baseline_ids = {
            str(block.get("id"))
            for block in (baseline.get(material) or {}).get("blocks") or []
            if isinstance(block, dict)
            and block.get("id")
            and block.get("content_floor", True)
        }
        touched = {
            baseline_id
            for baseline_id in baseline_ids
            if str((draft["baseline_dispositions"].get(baseline_id) or {}).get("action") or "retain") != "retain"
        }
        added = sum(
            1
            for block in draft[material]["blocks"]
            if str(block.get("change_action") or "") == "add"
        )
        total = len(baseline_ids)
        hard_limit = min(total, max(4, ceil(total * 0.60)))
        addition_limit = max(2, ceil(total * 0.10))
        if len(touched) > hard_limit:
            raise ValueError(
                f"baseline_transform_too_broad:{material}:{len(touched)}>{hard_limit}"
            )
        if added > addition_limit:
            raise ValueError(
                f"baseline_transform_too_many_additions:{material}:{added}>{addition_limit}"
            )
        ratio = (len(touched) / total) if total else 0.0
        scope[material] = {
            "baseline_blocks": total,
            "touched_baseline_blocks": len(touched),
            "added_blocks": added,
            "touched_ratio": round(ratio, 4),
            "review_level": "broad" if ratio > 0.35 else "focused",
        }
    draft["compiled_from"] = "bounded_baseline_transform"
    draft["transform_summary"] = {
        "rewritten_blocks": len(changed_ids),
        "added_blocks": len(additions),
        "retained_blocks": len(index) - len(changed_ids),
        "scope": scope,
        "transform_sha256": _digest(transform),
    }
    return draft


def build_tailoring_delta(baseline: dict[str, Any], draft: dict[str, Any]) -> dict[str, Any]:
    """Build the compact before/after review payload for the child auditor."""

    baseline_index: dict[str, dict[str, Any]] = {}
    baseline_positions: dict[str, int] = {}
    for material in ("cv", "cover_letter"):
        for position, block in enumerate((baseline.get(material) or {}).get("blocks") or []):
            if isinstance(block, dict) and block.get("id") and block.get("content_floor", True):
                baseline_index[str(block["id"])] = {**block, "material": material}
                baseline_positions[str(block["id"])] = position

    final_index: dict[str, tuple[str, int, dict[str, Any]]] = {}
    for material in ("cv", "cover_letter"):
        for position, block in enumerate((draft.get(material) or {}).get("blocks") or []):
            if isinstance(block, dict) and block.get("id"):
                final_index[str(block["id"])] = (material, position, block)

    dispositions = draft.get("baseline_dispositions") if isinstance(draft.get("baseline_dispositions"), dict) else {}
    changes: list[dict[str, Any]] = []
    seen_targets: set[str] = set()
    retained = 0
    for baseline_id, raw in dispositions.items():
        if not isinstance(raw, dict):
            continue
        action = str(raw.get("action") or "retain")
        if action == "retain":
            retained += 1
            continue
        target_id = str(raw.get("target_id") or "")
        if not target_id or target_id in seen_targets or target_id not in final_index:
            continue
        material, after_position, target = final_index[target_id]
        baseline_ids = [str(item) for item in (target.get("baseline_refs") or [baseline_id]) if str(item)]
        before_blocks = [baseline_index[item] for item in baseline_ids if item in baseline_index]
        changes.append(
            {
                "action": action,
                "material": material,
                "target_id": target_id,
                "baseline_ids": baseline_ids,
                "before": [str(item.get("text") or "") for item in before_blocks],
                "after": str(target.get("text") or ""),
                "before_positions": [baseline_positions[item] for item in baseline_ids if item in baseline_positions],
                "after_position": after_position,
                "section": str(target.get("section") or ""),
                "experience_id": str(target.get("experience_id") or ""),
                "priority": target.get("priority"),
                "jd_anchor_ids": list(target.get("jd_anchor_ids") or []),
            }
        )
        seen_targets.add(target_id)

    for material in ("cv", "cover_letter"):
        for position, block in enumerate((draft.get(material) or {}).get("blocks") or []):
            if not isinstance(block, dict) or str(block.get("change_action") or "") != "add":
                continue
            changes.append(
                {
                    "action": "add",
                    "material": material,
                    "target_id": str(block.get("id") or ""),
                    "baseline_ids": [],
                    "before": [],
                    "after": str(block.get("text") or ""),
                    "before_positions": [],
                    "after_position": position,
                    "section": str(block.get("section") or ""),
                    "experience_id": str(block.get("experience_id") or ""),
                    "priority": block.get("priority"),
                    "jd_anchor_ids": list(block.get("jd_anchor_ids") or []),
                }
            )
    return {
        "schema_version": 1,
        "baseline_sha256": str(baseline.get("baseline_sha256") or content_baseline_digest(baseline)),
        "mode": "bounded_incremental_transform",
        "changes": changes,
        "changed_block_count": len(changes),
        "retained_block_count": retained,
        "baseline_block_count": len(baseline_index),
        "deletion_allowed": False,
    }


def validate_content_floor(baseline: dict[str, Any], draft: dict[str, Any]) -> list[str]:
    """Validate traceable baseline coverage without judging prose similarity."""

    if not baseline:
        return ["content_baseline_missing"]
    errors: list[str] = []
    expected_digest = str(baseline.get("baseline_sha256") or content_baseline_digest(baseline))
    if str(draft.get("baseline_sha256") or "") != expected_digest:
        errors.append("baseline_sha256_mismatch")
    baseline_index: dict[str, tuple[str, dict[str, Any]]] = {}
    for material in ("cv", "cover_letter"):
        for block in (baseline.get(material) or {}).get("blocks") or []:
            if isinstance(block, dict) and block.get("id") and block.get("content_floor", True):
                baseline_index[str(block["id"])] = (material, block)
    final_index: dict[str, tuple[str, dict[str, Any]]] = {}
    reference_counts = {baseline_id: 0 for baseline_id in baseline_index}
    for material in ("cv", "cover_letter"):
        for block in (draft.get(material) or {}).get("blocks") or []:
            if not isinstance(block, dict) or not block.get("id"):
                continue
            block_id = str(block["id"])
            final_index[block_id] = (material, block)
            action = str(block.get("change_action") or "retain")
            if action in {"delete", "remove", "omit"}:
                errors.append(f"baseline_deletion_forbidden:{block_id}")
            for baseline_id in block.get("baseline_refs") or []:
                baseline_id = str(baseline_id)
                if baseline_id not in baseline_index:
                    errors.append(f"baseline_ref_unknown:{baseline_id}")
                else:
                    reference_counts[baseline_id] += 1
                    if baseline_index[baseline_id][0] != material:
                        errors.append(f"baseline_ref_material_mismatch:{baseline_id}")
    dispositions = draft.get("baseline_dispositions") if isinstance(draft.get("baseline_dispositions"), dict) else {}
    if set(dispositions) != set(baseline_index):
        for baseline_id in sorted(set(baseline_index) - set(dispositions)):
            errors.append(f"baseline_disposition_missing:{baseline_id}")
        for baseline_id in sorted(set(dispositions) - set(baseline_index)):
            errors.append(f"baseline_disposition_unknown:{baseline_id}")
    allowed = {"retain", "rewrite", "reorder", "merge"}
    for baseline_id, (material, _source) in baseline_index.items():
        disposition = dispositions.get(baseline_id)
        if not isinstance(disposition, dict):
            continue
        action = str(disposition.get("action") or "")
        if action not in allowed:
            errors.append(f"baseline_disposition_action_invalid:{baseline_id}")
        target_id = str(disposition.get("target_id") or "")
        target = final_index.get(target_id)
        if not target:
            errors.append(f"baseline_block_unrepresented:{baseline_id}")
            continue
        if target[0] != material or baseline_id not in (target[1].get("baseline_refs") or []):
            errors.append(f"baseline_block_unrepresented:{baseline_id}")
        if reference_counts.get(baseline_id) != 1:
            errors.append(f"baseline_reference_count_invalid:{baseline_id}")
    return sorted(set(errors))
